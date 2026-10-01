# -*- coding: utf-8 -*-
"""
合并所有补充表 → Supplementary_Tables_FINAL.xlsx
P0 原有 9 个 sheet + S6b 校准ECE + S8 成本(A实测/B微基准) + S9 λ敏感性
"""
import sys
from pathlib import Path
sys.stdout.reconfigure(encoding='utf-8')
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(r'd:\lht')
PR = ROOT / 'supplement_output' / 'paper_revision'
P0 = PR / 'Supplementary_Tables_P0.xlsx'
OUT = PR / 'Supplementary_Tables_FINAL.xlsx'

# ---- 读入各表 ----
p0 = pd.ExcelFile(P0)
s6b = pd.read_csv(PR / 'TableS6b_calibration_ece.csv')
s8a = pd.read_csv(PR / 'TableS8_cost.csv')
s8b = pd.read_csv(PR / 'TableS8_benchmark.csv')
s9 = pd.read_csv(PR / 'TableS9_lambda.csv')

# S9 列名中文化 + 标记基准行
s9 = s9.rename(columns={
    'tag': '实验配置', 'lambda1': 'λ₁(对比)', 'lambda2': 'λ₂(一致性)',
    'dataset': '数据集', 'epochs': 'Epochs', 'R2': 'R²', 'EF1': 'EF@1%', 'ECE': 'ECE'})
s9['是否论文基准'] = s9['实验配置'].apply(lambda x: '★基准' if x == 'base_l1_0.050_l2_0.10' else '')

HEAD_FILL = PatternFill('solid', fgColor='006D77')
HEAD_FONT = Font(bold=True, color='FFFFFF', size=10)
TITLE_FONT = Font(bold=True, color='006D77', size=11)
BASE_FILL = PatternFill('solid', fgColor='FFF2CC')


def style_sheet(ws, header_rows=(1,), ncols=None, freeze='A2'):
    ncols = ncols or ws.max_column
    for hr in header_rows:
        for c in range(1, ncols + 1):
            cell = ws.cell(row=hr, column=c)
            cell.fill = HEAD_FILL
            cell.font = HEAD_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    if freeze:
        ws.freeze_panes = freeze
    # 列宽自适应
    for c in range(1, ncols + 1):
        L = get_column_letter(c)
        width = 8
        for cell in ws[L]:
            if cell.value is not None:
                width = max(width, min(28, len(str(cell.value)) * 1.15 + 2))
        ws.column_dimensions[L].width = width


with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
    # 1) P0 原表按顺序复制
    for sheet in p0.sheet_names:
        pd.read_excel(P0, sheet_name=sheet).to_excel(xw, sheet_name=sheet, index=False)

    # 2) S6b 校准 ECE（Fig S6 配套）
    s6b.to_excel(xw, sheet_name='S6b_calibration', index=False)

    # 3) S9 λ 敏感性
    s9.to_excel(xw, sheet_name='S9_lambda', index=False)

    # 4) S8 成本：A 实测块 + B 微基准块，同一 sheet 上下排布
    s8a.to_excel(xw, sheet_name='S8_cost', index=False, startrow=1)
    wsa = xw.sheets['S8_cost']
    wsa.cell(row=1, column=1, value='块A：真实训练日志实测（RTX 4090，AMP；— 表示当年控制台日志未留存）').font = TITLE_FONT
    # B 块放在 A 块下方空 2 行
    b_start = wsa.max_row + 3
    s8b.to_excel(xw, sheet_name='S8_cost', index=False, startrow=b_start)
    wsa.cell(row=b_start, column=1,
             value='块B：统一微基准（RTX 4090，AMP，合成batch，前向+反向×10，仅横向相对比较）').font = TITLE_FONT
    ncol_b = s8b.shape[1]
    for c in range(1, ncol_b + 1):
        cell = wsa.cell(row=b_start + 1, column=c)
        cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # ---- 统一样式 ----
    wb = xw.book
    for name in p0.sheet_names + ['S6b_calibration', 'S9_lambda']:
        style_sheet(wb[name])
    # S8：A 块标题行=2，冻结到 A3
    style_sheet(wb['S8_cost'], header_rows=(2,), freeze='A3')
    # S8 B 块表头样式已单独处理；列宽
    for c in range(1, wsa.max_column + 1):
        L = get_column_letter(c)
        width = 8
        for cell in wsa[L]:
            if cell.value is not None:
                width = max(width, min(30, len(str(cell.value)) * 1.1 + 2))
        wsa.column_dimensions[L].width = width
    # S9 基准行高亮
    ws9 = wb['S9_lambda']
    for r in range(2, ws9.max_row + 1):
        if ws9.cell(row=r, column=ws9.max_column).value == '★基准':
            for c in range(1, ws9.max_column + 1):
                ws9.cell(row=r, column=c).fill = BASE_FILL

print(f'✅ 最终工作簿已保存 → {OUT}')
xl = pd.ExcelFile(OUT)
print('Sheets 顺序:')
for i, s in enumerate(xl.sheet_names, 1):
    d = pd.read_excel(OUT, sheet_name=s)
    print(f'  {i:2d}. {s:22s} {d.shape[0]} 行 × {d.shape[1]} 列')
