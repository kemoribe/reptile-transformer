# -*- coding: utf-8 -*-
"""
补充任务3：汇总 7 个模型在 4 个数据集上的 EF@1% / EF@5%，输出规律统计。
数据来源：
  现象表   d:\\lht\\实验结果\\ChEMBL_Davis模型性能对比（现象）.xlsx （6 个模型）
  消融表   d:\\lht\\实验结果\\ChEMBL_Davis模型性能对比（内部+外部消融对比实验）.xlsx
           （补充 Reptile+Transformer 行；其余 6 模型两表数值一致）
"""
import sys
from pathlib import Path
import pandas as pd
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

F_PHENO = Path(r'd:\lht\实验结果\ChEMBL_Davis模型性能对比（现象）.xlsx')
F_ABLAT = Path(r'd:\lht\实验结果\ChEMBL_Davis模型性能对比（内部+外部消融对比实验）.xlsx')
OUT = Path(r'd:\lht\supplement_output')
OUT.mkdir(parents=True, exist_ok=True)

SHEET2DS = {'ChEMBL数据集': 'ChEMBL', 'Davis数据集': 'Davis',
            'BingdingBD数据集': 'BindingDB', 'KIBA数据集': 'KIBA'}

REPTILE = 'Reptile+Transformer'
TRANS = 'Transformer'
MLP = 'MLP'
GNNS = ['GraphDTA GAT_GCN', 'GraphDTA GCNNet', 'GraphDTA GATNet', 'GraphDTA GINConnvNet']
ORDER = [REPTILE, TRANS, MLP] + GNNS


def norm_model(name):
    if name is None:
        return None
    s = str(name)
    if s.startswith('Reptile'):
        return REPTILE
    if s.startswith('Transformer'):
        return TRANS
    if s.startswith('MLP'):
        return MLP
    for g in GNNS:
        if g in s:
            return g
    return None


def read_table(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    for sheet, ds in SHEET2DS.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for r in ws.iter_rows(min_row=2, values_only=True):
            m = norm_model(r[0])
            if m is None or r[6] is None:
                continue
            rows.append({'dataset': ds, 'model': m,
                         'R2': float(r[1]) if isinstance(r[1], (int, float)) else None,
                         'EF1': float(r[6]), 'EF5': float(r[7]), 'EF10': float(r[8]),
                         'source': path.stem})
    return pd.DataFrame(rows)


def main():
    pheno = read_table(F_PHENO)
    ablat = read_table(F_ABLAT)

    # 现象表无 Reptile，从消融表补；其余取现象表（两表一致）
    rep = ablat[ablat.model == REPTILE]
    df = pd.concat([pheno, rep], ignore_index=True)
    df = df.drop_duplicates(subset=['dataset', 'model'], keep='first')
    df['model'] = pd.Categorical(df['model'], categories=ORDER, ordered=True)
    df = df.sort_values(['dataset', 'model'])

    ef1 = df.pivot(index='model', columns='dataset', values='EF1').reindex(ORDER)[
        ['ChEMBL', 'Davis', 'BindingDB', 'KIBA']]
    ef5 = df.pivot(index='model', columns='dataset', values='EF5').reindex(ORDER)[
        ['ChEMBL', 'Davis', 'BindingDB', 'KIBA']]

    print('=' * 78)
    print('EF@1%（每数据集；GNN Davis 数值为论文逐靶点均值，小样本噪声大）')
    print('=' * 78)
    print(ef1.round(3).to_string())
    print('\n' + '=' * 78)
    print('EF@5%')
    print('=' * 78)
    print(ef5.round(3).to_string())

    # 排名（1=最高富集）
    r1 = ef1.rank(ascending=False, axis=0).mean(axis=1)
    r5 = ef5.rank(ascending=False, axis=0).mean(axis=1)
    tab = pd.DataFrame({
        'EF1_mean': ef1.mean(axis=1).round(3),
        'EF1_avgRank': r1.round(2),
        'EF5_mean': ef5.mean(axis=1).round(3),
        'EF5_avgRank': r5.round(2),
        'EF1_minus_EF5': (ef1 - ef5).mean(axis=1).round(3),
    })
    print('\n' + '=' * 78)
    print('跨四数据集均值与平均排名')
    print('=' * 78)
    print(tab.to_string())

    # 非 GNN vs GNN 分组（Davis GNN 小样本失真，同时给排除 Davis 的均值）
    nongnn = [REPTILE, TRANS, MLP]
    grp = pd.DataFrame({
        'EF1_all4': ef1.mean(axis=1),
        'EF5_all4': ef5.mean(axis=1),
        'EF1_exclDavis': ef1[['ChEMBL', 'BindingDB', 'KIBA']].mean(axis=1),
        'EF5_exclDavis': ef5[['ChEMBL', 'BindingDB', 'KIBA']].mean(axis=1),
    })
    grp['group'] = ['非GNN', '非GNN', '非GNN', 'GNN', 'GNN', 'GNN', 'GNN']
    print('\n' + '=' * 78)
    print('组均值')
    print('=' * 78)
    print(grp.groupby('group').mean().round(3).to_string())

    df.to_csv(OUT / 'ef7x4_long.csv', index=False, encoding='utf-8-sig')
    with pd.ExcelWriter(OUT / 'EF规律汇总.xlsx') as w:
        ef1.to_excel(w, sheet_name='EF@1%')
        ef5.to_excel(w, sheet_name='EF@5%')
        tab.to_excel(w, sheet_name='均值与排名')
        grp.groupby('group').mean().to_excel(w, sheet_name='分组均值')
    print('\n已保存:', OUT / 'EF规律汇总.xlsx')


if __name__ == '__main__':
    main()
