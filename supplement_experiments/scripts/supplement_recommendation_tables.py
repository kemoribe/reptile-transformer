# -*- coding: utf-8 -*-
"""
补充实验整理：
  表1  四数据集：跨簇比例 / GNN_EF1均值 / 非GNN_EF1均值 / ΔEF1 / 推荐模型
  图1  跨簇比例 vs EF@1% 均值（GNN vs 非GNN）
  表2  Tanimoto 相似度阈值 0.10/0.13/0.16 下，R² 准则 vs EF@1% 准则的推荐模型对比
输出：d:\\lht\\supplement_output\\
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path

import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from scipy.stats import spearmanr
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = Path(r'd:\lht\supplement_output')

# ---------------------------------------------------------------- 数据
# EF@1%（现象表6模型 + 消融表 Reptile 行）
EF1 = {
    'Reptile':    {'chembl': 1.98, 'davis': 1.75, 'kiba': 4.72, 'bindingdb': 1.85},
    'Transformer': {'chembl': 4.44, 'davis': 4.20, 'kiba': 4.05, 'bindingdb': 4.68},
    'MLP':        {'chembl': 4.63, 'davis': 4.20, 'kiba': 3.75, 'bindingdb': 3.89},
    'GAT_GCN':    {'chembl': 2.11, 'davis': 0.4276, 'kiba': 4.49, 'bindingdb': 3.02},
    'GCNNet':     {'chembl': 2.41, 'davis': 0.4315, 'kiba': 4.28, 'bindingdb': 3.37},
    'GATNet':     {'chembl': 1.59, 'davis': 0.469, 'kiba': 3.88, 'bindingdb': 2.46},
    'GINConvNet': {'chembl': 2.84, 'davis': 4.34, 'kiba': 4.37, 'bindingdb': 3.45},
}
R2 = {
    'Reptile':    {'chembl': 0.202, 'davis': 0.2237, 'kiba': 0.3222, 'bindingdb': 0.2335},
    'Transformer': {'chembl': 0.1679, 'davis': 0.1713, 'kiba': 0.2308, 'bindingdb': 0.1838},
    'MLP':        {'chembl': 0.1806, 'davis': 0.1713, 'kiba': 0.2565, 'bindingdb': 0.2220},
    'GAT_GCN':    {'chembl': 0.0579, 'davis': 0.3998, 'kiba': 0.2233, 'bindingdb': 0.0995},
    'GCNNet':     {'chembl': 0.0447, 'davis': 0.3660, 'kiba': 0.2122, 'bindingdb': 0.0519},
    'GATNet':     {'chembl': -0.1528, 'davis': 0.3667, 'kiba': -0.1038, 'bindingdb': 0.0891},
    'GINConvNet': {'chembl': -0.0913, 'davis': 0.2146, 'kiba': 0.1203, 'bindingdb': 0.0677},
}
GNN = ['GAT_GCN', 'GCNNet', 'GATNet', 'GINConvNet']
NON = ['Reptile', 'Transformer', 'MLP']

DSS = ['chembl', 'davis', 'kiba', 'bindingdb']
DS_CN = {'chembl': 'ChEMBL', 'davis': 'Davis', 'kiba': 'KIBA', 'bindingdb': 'BindingDB'}
# 跨簇比例（MMseqs2, 序列一致性40%, 覆盖度0.8）
CROSS = {'chembl': 53.85, 'davis': 41.54, 'kiba': 41.18, 'bindingdb': 53.57}
# 训练集药物平均 Tanimoto（analyze_tanimoto.py，ECFP4）
TANI = {'chembl': 0.1221, 'davis': 0.1367, 'kiba': 0.2500, 'bindingdb': 0.3600}

gnn_mean = {d: float(np.mean([EF1[m][d] for m in GNN])) for d in DSS}
non_mean = {d: float(np.mean([EF1[m][d] for m in NON])) for d in DSS}
delta = {d: non_mean[d] - gnn_mean[d] for d in DSS}

best_ef1 = {d: max(list(EF1), key=lambda m: EF1[m][d]) for d in DSS}
best_r2 = {d: max(list(R2), key=lambda m: R2[m][d]) for d in DSS}

# ---------------------------------------------------------------- 图1
plt.rcParams.update({
    'font.family': 'Arial', 'axes.linewidth': 1.1,
    'xtick.labelsize': 10.5, 'ytick.labelsize': 10.5,
    'axes.labelsize': 12, 'legend.fontsize': 10,
})
TEAL_D, TEAL_L = '#006D77', '#83C5BE'
GRAY = '#9AA5A8'

fig, ax = plt.subplots(figsize=(8.2, 5.4), dpi=300)

x = np.array([CROSS[d] for d in DSS])
yg = np.array([gnn_mean[d] for d in DSS])
yn = np.array([non_mean[d] for d in DSS])

# 同数据集配对连线
for xi, a, b in zip(x, yg, yn):
    ax.plot([xi, xi], [a, b], color=GRAY, ls=':', lw=1.0, zorder=1)

ax.axhline(1.0, color='#B0B0B0', ls='--', lw=1.0, zorder=0)
ax.text(57.6, 1.06, 'random enrichment = 1', fontsize=8.5, color='#777',
        ha='right', va='bottom')

ax.scatter(x, yg, s=170, c=TEAL_L, edgecolors=TEAL_D, linewidths=1.4,
           marker='o', zorder=3, label='GNN mean (4 GraphDTA models)')
ax.scatter(x, yn, s=200, c=TEAL_D, edgecolors='white', linewidths=1.2,
           marker='s', zorder=3, label='Non-GNN mean (Reptile / Transformer / MLP)')

# 每个数据集一个标签，引线指向配对中部（ChEMBL/BindingDB、Davis/KIBA 的 x 极近，错位避让）
label_pos = {
    'chembl':    (55.6, 3.95),
    'davis':     (38.4, 1.30),
    'kiba':      (43.4, 3.80),
    'bindingdb': (48.2, 3.30),
}
for d in DSS:
    xm, ym = CROSS[d], (gnn_mean[d] + non_mean[d]) / 2
    tx, ty = label_pos[d]
    ax.annotate(DS_CN[d], xy=(xm, ym), xytext=(tx, ty),
                fontsize=10.5, fontweight='bold', color='#1B3A4B',
                arrowprops=dict(arrowstyle='-', color='#90A4AE', lw=0.8),
                ha='left' if tx > xm else 'left')

rho_g = spearmanr(x, yg).statistic
rho_n = spearmanr(x, yn).statistic
ax.set_xlabel('Cross-cluster test-target ratio (%)\n(MMseqs2, seq identity $\\geq$ 40%, coverage $\\geq$ 0.8)')
ax.set_ylabel('Mean EF@1%')
ax.set_xlim(37, 60.2)
ax.set_ylim(0.6, 5.2)
ax.grid(color='#E4ECEB', ls='--', lw=0.7, zorder=0)
ax.spines['top'].set_visible(False)
ax.spines['right'].set_visible(False)
ax.legend(loc='upper left', frameon=True, framealpha=0.95, edgecolor='#CFD8DC')

# Spearman 标注放图外（右侧留白）
fig.subplots_adjust(right=0.80)
fig.text(0.825, 0.62,
         f'Spearman ρ (n = 4)\n'
         f'GNN: ρ = {rho_g:+.2f}\n'
         f'Non-GNN: ρ = {rho_n:+.2f}',
         fontsize=9.5, va='top', ha='left',
         bbox=dict(boxstyle='round,pad=0.5', fc='#F2F8F7', ec=TEAL_D, lw=0.9))

for ext in ('png', 'tiff'):
    kw = {'dpi': 300, 'bbox_inches': 'tight', 'facecolor': 'white'}
    if ext == 'tiff':
        kw['pil_kwargs'] = {'compression': 'tiff_lzw'}
    fig.savefig(OUT / f'fig1_crosscluster_vs_ef1.{ext}', **kw)
plt.close(fig)
print('图1 已保存')

# ---------------------------------------------------------------- Excel
wb = Workbook()
thin = Side(style='thin', color='B0BEC5')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill('solid', fgColor='006D77')
sub_fill = PatternFill('solid', fgColor='E3F0EE')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)


def style_sheet(ws, ncol, widths):
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ncol):
        for c in row:
            c.border = border
            if c.row == 1:
                c.font = Font(name='Arial', bold=True, color='FFFFFF', size=10.5)
                c.fill = hdr_fill
                c.alignment = center
            else:
                c.font = Font(name='Arial', size=10.5)
                c.alignment = center if c.column != ncol else left

# ---------- 表1
ws1 = wb.active
ws1.title = '表1_跨簇比例与推荐'
ws1.append(['数据集', '跨簇比例 (%)\n(一致性40%)', 'GNN\nEF@1%均值', '非GNN\nEF@1%均值',
            'ΔEF1\n(非GNN−GNN)', '推荐模型（虚拟筛选EF导向）'])
rec = {
    'chembl': 'MLP（EF@1%=4.63，最高且ECE低）',
    'davis': 'MLP / Transformer（EF@1%=4.20，稳健；GINConvNet名义4.34但Davis小样本失真）',
    'kiba': 'Reptile+Transformer（EF@1%=4.72 且 R²=0.322，双优）',
    'bindingdb': 'Transformer（EF@1%=4.68）',
}
for d in DSS:
    ws1.append([DS_CN[d], CROSS[d], round(gnn_mean[d], 2), round(non_mean[d], 2),
                f'{delta[d]:+.2f}', rec[d]])
notes1 = [
    '注：ΔEF1 = 非GNN组（Reptile+Transformer / Transformer / MLP）EF@1%均值 − GNN组（GAT_GCN/GCNNet/GATNet/GINConvNet）均值；正值表示非GNN占优。',
    '注：跨簇比例为测试靶点中与全部训练靶点序列一致性低于40%（MMseqs2，覆盖度≥0.8）者的比例；60%/80%阈值下比例更高（详见protein_cluster/cross_cluster_summary.csv）。',
    '注：Davis每靶点仅66个化合物，EF@1%每靶点仅取1个，GNN的EF值（0.43/4.34）噪声大，解读需谨慎。',
]
for n in notes1:
    ws1.append([n])
style_sheet(ws1, 6, [12, 14, 11, 11, 13, 52])
for r in range(2, 6):
    ws1.row_dimensions[r].height = 24
ws1.row_dimensions[1].height = 34
for r in range(7, 10):
    ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws1.cell(r, 1).font = Font(name='Arial', size=9, italic=True, color='555555')
    ws1.cell(r, 1).alignment = left

# ---------- 表2
ws2 = wb.create_sheet('表2_相似度阈值推荐')
ws2.append(['Tanimoto\n阈值 τ', '相似度区间', '落入数据集\n(平均Tanimoto)',
            'R²准则推荐（模型:R²）', 'EF@1%准则推荐（模型:EF@1%）', '两准则一致性'])


def winners(dss):
    r2_txt = '；'.join(f"{DS_CN[d]}→{best_r2[d]}({R2[best_r2[d]][d]:.3f})" for d in dss)
    ef_txt = '；'.join(f"{DS_CN[d]}→{best_ef1[d]}({EF1[best_ef1[d]][d]:.2f})" for d in dss)
    agree = [d for d in dss if best_r2[d] == best_ef1[d]]
    if not dss:
        return '—', '—', '—'
    ag_txt = f"{len(agree)}/{len(dss)} 一致" + (f"（{'、'.join(DS_CN[d] for d in agree)}）" if agree else '（无）')
    return r2_txt, ef_txt, ag_txt


for tau in (0.10, 0.13, 0.16):
    low = [d for d in DSS if TANI[d] < tau]
    high = [d for d in DSS if TANI[d] >= tau]
    for regime, grp in (('低相似度 (<τ)', low), ('高相似度 (≥τ)', high)):
        if grp:
            r2t, eft, ag = winners(grp)
            dtxt = '、'.join(f"{DS_CN[d]}({TANI[d]:.3f})" for d in grp)
        else:
            r2t = eft = ag = '—'
            dtxt = '无（全部数据集 ≥ τ）'
        ws2.append([f'{tau:.2f}' if regime.startswith('低') else '', regime, dtxt, r2t, eft, ag])

notes2 = [
    '注：阈值τ针对训练集药物平均ECFP4 Tanimoto相似度（ChEMBL 0.122、Davis 0.137、KIBA 0.250、BindingDB 0.360），用于划分低/高化学空间相似度场景。',
    '注：各数据集在7个模型中分别取R²最优与EF@1%最优者；两准则仅在KIBA上一致（Reptile+Transformer），ChEMBL/Davis/BindingDB均出现"R²优但富集弱"（Reptile/GAT_GCN）与"富集优"（MLP/Transformer）的分化。',
    '注：Davis的GNN EF@1%受每靶点66化合物的小样本影响（GINConvNet 4.34名义最高），实际筛选推荐MLP/Transformer（4.20）。',
]
for n in notes2:
    ws2.append([n])
style_sheet(ws2, 6, [10, 14, 30, 44, 44, 18])
ws2.row_dimensions[1].height = 34
for r in range(2, ws2.max_row + 1):
    ws2.row_dimensions[r].height = 40
# 阈值单元格合并 & 区间底色
for r0 in (2, 4, 6):
    ws2.merge_cells(start_row=r0, start_column=1, end_row=r0 + 1, end_column=1)
    ws2.cell(r0, 1).alignment = center
for r in (2, 4, 6):
    for c in range(1, 7):
        ws2.cell(r, c).fill = sub_fill
nr = ws2.max_row
for k, r in enumerate(range(nr - 2, nr + 1)):
    ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws2.cell(r, 1).font = Font(name='Arial', size=9, italic=True, color='555555')
    ws2.cell(r, 1).alignment = left

xlsx_path = OUT / '模型推荐与阈值敏感性_表1表2.xlsx'
wb.save(xlsx_path)
print('Excel 已保存:', xlsx_path)

# 控制台核对
print('\n表1:')
for d in DSS:
    print(f"  {DS_CN[d]:9s} cross={CROSS[d]:5.2f}%  GNN={gnn_mean[d]:.3f}  "
          f"non={non_mean[d]:.3f}  Δ={delta[d]:+.3f}  EF1winner={best_ef1[d]}  R2winner={best_r2[d]}")
print(f"\nSpearman: GNN {rho_g:+.3f}, non-GNN {rho_n:+.3f}")
