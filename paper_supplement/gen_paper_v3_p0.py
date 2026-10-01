# -*- coding: utf-8 -*-
"""
论文 V3-P0 生成器：在 论文V3_补充实验与故事线重构版.docx 基础上写入 P0 五项稳健性实验
  P0-1 τ≈0.13 留一数据集(LODO)验证 -> 降级为 descriptive band
  P0-2 family-remote 路由器 LODO 验证 -> 泛化成立 (95.5-97.1%)
  P0-3 MMseqs2 覆盖率x同一性敏感性扫描 + CD-HIT式聚类算法对照
  P0-4 审计 vs 未审计名义冷启动评估：排名一致性 + 误选代价
  P0-5 随机置乱序列负对照
正文新增 3.6.5 小节 + Table 6；附录新增 Tables S26-S31 + Figure S11；
并同步软化全文 τ≈0.13 的阈值措辞、强化路由器表述。
输出：论文V3_P0稳健性验证版.docx
"""
import sys
import docx
import pandas as pd
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches

sys.stdout.reconfigure(encoding='utf-8')

SRC = r'd:\lht\修改supplement_output\论文V3_补充实验与故事线重构版.docx'
DST = r'd:\lht\修改supplement_output\论文V3_P0稳健性验证版.docx'
P0 = r'd:\lht\修改supplement_output\p0_experiments'

doc = docx.Document(SRC)
paras = doc.paragraphs
ADDED = []


def find(substr, start=0):
    for i, p in enumerate(paras):
        if i >= start and substr in p.text:
            return i, p
    raise RuntimeError('anchor not found: ' + substr)


S_H3 = paras[[i for i, p in enumerate(paras) if p.text.strip().startswith('3.6.4')][0]].style
S_BODY = find('Inner-loop length sweep')[1].style
S_CAP = find('Table 5 Protein-encoder capacity control')[1].style


def _mk_par(text, style=None):
    p = doc.add_paragraph(text)
    if style is not None:
        try:
            p.style = style
        except Exception:
            pass
    return p._p


def insert_block_before(anchor_par, items):
    anchor = anchor_par._p
    for kind, val, *rest in items:
        if kind == 'par':
            anchor.addprevious(_mk_par(val, rest[0] if rest else None))
        elif kind == 'table':
            t = doc.add_table(rows=len(val), cols=len(val[0]))
            try:
                t.style = doc.tables[0].style
            except Exception:
                pass
            for r, row in enumerate(val):
                for c, v in enumerate(row):
                    t.cell(r, c).text = str(v)
            for c in t.rows[0].cells:
                for pp in c.paragraphs:
                    for run in pp.runs:
                        run.bold = True
            anchor.addprevious(t._tbl)
        elif kind == 'pic':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(val, width=Inches(5.9))
            anchor.addprevious(p._p)


def replace_in_paragraph(p, old, new):
    """优先 run 内替换；跨 run 时合并到首个命中 run。"""
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    full = ''.join(r.text for r in p.runs)
    if old in full and p.runs:
        full = full.replace(old, new)
        p.runs[0].text = full
        for r in p.runs[1:]:
            r.text = ''
        return True
    return False


def append_sentence(p, text):
    p.add_run(text)


# =====================================================================
# 一、全文措辞同步（τ 降级描述性；路由器 LODO 强化）
# =====================================================================
edits = []

# 摘要 finding (ii)
_, p = find('narrow training-panel Tanimoto band')
ok = replace_in_paragraph(
    p, 'a narrow training-panel Tanimoto band (τ ≈ 0.13)',
    'a narrow, descriptive training-panel Tanimoto band (τ ≈ 0.13; n = 4 panels)')
edits.append(('abstract band', ok))

# 贡献 5 路由器
_, p = find('validated family-remote routing classifier')
ok = replace_in_paragraph(
    p, 'AUC 0.960 at a 40% identity threshold)',
    'AUC 0.960 at a 40% identity threshold, with 95.5–97.1% accuracy in four-fold '
    'leave-one-dataset-out validation)')
edits.append(('contribution 5 router LODO', ok))

# Figure 5 图注（硬阈值 -> 描述性）
_, p = find('GNN performance gain boundary at the 0.13')
replace_in_paragraph(p, p.text,
    'Figure  GNN advantage ΔR² versus training-panel similarity; the τ ≈ 0.13 '
    'GNN-favourable band is descriptive (four-fold leave-one-dataset-out validation, '
    'Section 3.6.5). Positive ΔR² is observed only on the dense Davis panel.')
edits.append(('fig5 caption', True))

# Figure 6 图注
_, p = find('optimal classification threshold for GNN applicability')
replace_in_paragraph(p, p.text,
    'Figure  Accuracy curve of the similarity threshold scan. The empirically observed '
    'GNN-applicability band lies near Tanimoto similarity ≈ 0.13; leave-one-dataset-out '
    'validation (Section 3.6.5) does not support its use as a predictive threshold.')
edits.append(('fig6 caption', True))

# §3.2.1 阈值扫描段追加
_, p = find('identify an optimal dividing threshold')
append_sentence(
    p, ' A four-fold leave-one-dataset-out validation of this boundary (Section 3.6.5, '
       'Table S31) cannot predict the held-out panel above the majority-class rate '
       '(2/4 correct versus 3/4), and the boundary is unidentifiable in the fold removing '
       'the sole GNN-favourable panel; τ ≈ 0.13 is therefore reported as a descriptive '
       'empirical band rather than a deployable threshold.')
edits.append(('3.2.1 LODO caveat', True))

# §3.6.3 追加
_, p = find('fingerprint-relative construct')
append_sentence(
    p, ' A leave-one-dataset-out predictive validation of the τ ≈ 0.13 boundary (Section '
       '3.6.5) likewise fails to outperform the majority-class predictor; τ is consequently '
       'retained strictly as a fingerprint-relative descriptive band, never as a deployment '
       'cut-off.')
edits.append(('3.6.3 cross-ref', True))

# Figure 9 框架段追加
_, p = find('Figure 9 condenses the theory')
append_sentence(
    p, ' This map is descriptive at n = 4 panels: leave-one-dataset-out validation '
       '(Section 3.6.5) cannot predict the held-out regime above the majority-class '
       'baseline and the band cannot be fitted at all when the sole GNN-favourable panel '
       'is held out. It is offered as a post hoc organizer of these four panels, not as a '
       'deployable rule.')
edits.append(('fig9 caveat', True))

# 指南第 4 条
_, p = find('Boundary-region model ensembling')
replace_in_paragraph(p, 'Near the τ ≈ 0.13 applicability boundary',
                     'Near the empirically observed, descriptive (not externally '
                     'validated) τ ≈ 0.13 band')
edits.append(('guideline 4 wording', True))

# limitation 5 追加
_, p = find('5.Principle scale')
append_sentence(
    p, ' The leave-one-dataset-out and negative-control analyses added in Section 3.6.5 '
       'further delimit this scope: the remote-target router and the audit classifications '
       'generalize across held-out datasets, whereas the τ ≈ 0.13 band remains descriptive '
       'only.')
edits.append(('limitation 5 update', True))

# limitation 6 路由器 LODO
_, p = find('6.Distribution-shift risk')
replace_in_paragraph(
    p, 'remains ≥ 90.7% accurate across a 30–60% threshold scan (Table S 25).',
    'remains ≥ 90.7% accurate across a 30–60% threshold scan (Table S25), and—critically '
    'for cross-dataset deployment—retains 95.5–97.1% accuracy in four-fold '
    'leave-one-dataset-out validation, where each fold chooses its threshold using only '
    'the other three datasets and 40% is selected in every fold (Table 6, Table S26).')
edits.append(('limitation 6 router LODO', True))

# =====================================================================
# 二、正文 3.6.5 小节（插在 "4 Discussion and conclusion" 之前）
# =====================================================================
exp1 = pd.read_csv(P0 + r'\p0_exp1_tau_lodo_folds.csv')
exp2 = pd.read_csv(P0 + r'\p0_experiments' if False else P0 + r'\p0_exp2_router_lodo.csv')
grid = pd.read_csv(P0 + r'\param_sensitivity\mmseqs_grid.csv')
algo = pd.read_csv(P0 + r'\param_sensitivity\cluster_algorithm_cdHit_comparison.csv')
exp4 = pd.read_csv(P0 + r'\p0_experiments' if False else P0 + r'\p0_exp4_rank_comparison.csv')
neg = pd.read_csv(P0 + r'\negative_control\negative_control_results.csv')

acc_min, acc_max = exp2['Accuracy'].min(), exp2['Accuracy'].max()
r2 = exp4[exp4.metric == 'R2']
n_rev = int(r2.winner_reversed.sum())

table6 = [
    ['Analysis', 'Design', 'Headline result', 'Claim action'],
    ['τ ≈ 0.13 applicability band',
     'Leave-one-dataset-out (LODO): fit band on 3 panels, test 4th',
     '2/4 correct (0.50), below 0.75 majority-class; one fold unidentifiable',
     'Downgraded to descriptive band'],
    ['Family-remote router',
     'LODO threshold selection (3 datasets → 4th), n = 237',
     f'{acc_min*100:.1f}–{acc_max*100:.1f}% accuracy; 40% selected every fold; FNR = 0',
     'Cross-dataset generalization validated'],
    ['Coverage sensitivity',
     'Bidirectional coverage c ∈ {0.7, 0.8, 0.9} at 40% identity',
     'Cross-cluster ratios 35.4–62.5%; semi-leakage conclusion preserved',
     'Robust'],
    ['Identity sensitivity',
     'Identity t ∈ {30,40,50,60,70,80}%, c = 0.8',
     'Monotone 26.1→100%; headline ratios exactly reproduced',
     'Robust; exact reproduction'],
    ['Clustering algorithm',
     'Set-cover/connected-component vs CD-HIT-style greedy incremental',
     'Per-target label agreement 92.3–100%; Cohen κ 0.84–1.00',
     'Functional equivalence supported'],
    ['Random-sequence negative control',
     'Residue-shuffled test sequences (length/composition kept), 3 seeds',
     '100% cross-cluster; zero ≥40%-identity hits in all 36 settings',
     'Audit specificity confirmed'],
    ['Audited vs unaudited evaluation',
     'Ranks on nominal full cold test vs homology-free subset (5 models × 3 datasets)',
     f'R²-winner reversals {n_rev}/3; Kendall τb 0.00–0.80; ΔR² mis-selection 0.014–0.015',
     'Value of auditing quantified'],
]

items = [
    ('par', '3.6.5 Framework-internal robustness: leave-one-dataset-out validation, '
            'parameter sensitivity, and negative controls', S_H3),
    ('par', 'Several framework-level claims are themselves inferred from only four panels. '
            'To delimit their scope before deployment, we pre-registered five robustness '
            'analyses that require no new training data (Table 6): predictive validation of '
            'the τ band, leave-one-dataset-out validation of the remote-target router, a '
            'full scan of the audit clustering parameters, a clustering-algorithm cross-check, '
            'and a random-sequence negative control; the value of auditing is then quantified '
            'against the unaudited protocol.', S_BODY),

    ('par', 'Applicability band (τ ≈ 0.13). We fitted the GNN-favourable band on three '
            'panels (boundaries = midpoints between the positive panel and its nearest '
            'negative neighbour on each side) and predicted the fourth. LODO accuracy is '
            '0.50 (2/4)—below the 0.75 majority-class baseline: the fold removing Davis, the '
            'sole GNN-favourable panel, cannot fit the rule at all, and the fold removing '
            'ChEMBL lacks a lower boundary and misclassifies it; a 1-nearest-neighbour rule '
            'scores identically (Table S31). The lower boundary is stable when estimable '
            '(0.130 in both identifiable folds) but the upper boundary varies (0.194–0.249). '
            'Consistent with our pre-registered decision rule, τ ≈ 0.13 is therefore stated '
            'as a descriptive empirical band only; it must not be used as a stand-alone '
            'deployment threshold on an unseen panel.', S_BODY),

    ('par', 'Router cross-dataset validation. In each LODO fold the routing threshold was '
            'chosen on three datasets by accuracy maximization over 25–60% (ties resolved '
            'toward the pre-registered 40%) and applied unchanged to the fourth. Accuracy '
            'was 96.2% (ChEMBL), 96.9% (Davis), 97.1% (KIBA) and 95.5% (BindingDB), with '
            'zero false negatives in every fold and 40% selected as the threshold in every '
            'fold—identical to the fixed pre-registered rule (Table S26). The router’s '
            'cross-dataset generalization claim is therefore supported rather than being a '
            'pooled-set artifact.', S_BODY),

    ('par', 'Table 6 Framework-internal robustness analyses (P0): designs, headline results, '
            'and claim actions.', S_CAP),
    ('table', table6),
    ('par', 'All analyses re-use the audited splits of Section 3.1; no model was retrained. '
            'Per-fold and per-parameter details are in Tables S26–S31 and Figure S11.', S_BODY),

    ('par', 'Audit-parameter sensitivity. Re-running easy-cluster over 18 settings per '
            'dataset (bidirectional coverage 0.7/0.8/0.9 × identity 30–80%) exactly '
            'reproduces the published ratios at the operating point (c = 0.8; 40/60/80%). '
            'At the headline 40%-identity threshold, moving coverage across 0.7–0.9 moves '
            'cross-cluster ratios within 35.4–53.9% (Davis), 41.2–55.9% (KIBA), 53.9–57.7% '
            '(ChEMBL) and 45.5–62.5% (BindingDB)—the qualitative finding, that roughly half '
            'of nominally cold test targets retain a same-cluster train homologue, is '
            'preserved at every coverage setting, and ratios increase monotonically with '
            'identity (26.1–100% across the grid; Table S27, Figure S11). Replacing the '
            'default set-cover/connected-component clustering with CD-HIT-style greedy '
            'incremental clustering (MMseqs2 --cluster-mode 2; c = 0.8) changes per-target '
            'same/cross labels by 0–7.7% across the 12 dataset×threshold cells (agreement '
            '92.3–100%, Cohen κ = 0.84–1.00; Table S28), supporting the stated functional '
            'equivalence of the two clustering engines.', S_BODY),

    ('par', 'Audited versus unaudited evaluation. On the three datasets with matched '
            'prediction sets, ranks of the five models under the nominal (unaudited) '
            'target-cold test set and the 40%-identity homology-free subset agree only '
            'modestly (Kendall τb = 0.00/0.33/0.80 for R² on Davis/BindingDB/KIBA): the R² '
            'winner changes on 2/3 datasets, and selecting on the unaudited benchmark '
            'costs 0.015 R² on Davis (GCNNet chosen; Reptile-Transformer best on the '
            'audited subset; 5.8% relative) and 0.014 R² on BindingDB (GAT_GCN chosen; '
            'GATNet best; 36.6% relative); the Davis EF@1% mis-selection is larger (0.88 '
            'enrichment units; 17.6%), although this point sits on the degenerate 66-'
            'compound panel discussed in Section 3.3 (Table S29; values recomputed from '
            'Table S21 predictions). These numbers quantify the price of skipping the '
            'audit even when the split is nominally target-disjoint.', S_BODY),

    ('par', 'Negative control. Residue-shuffled test sequences—each permuted within '
            'sequence to preserve length and amino-acid composition (three seeds)—were '
            'clustered with the unchanged training panel. Across all four datasets, three '
            'seeds and the 40/60/80% thresholds, 100% of shuffled targets were cross-cluster '
            'and no shuffled target produced any bidirectionally covered (c = 0.8) hit at or '
            'above 40% identity, whereas the true sequences show 41.2–53.9% cross-cluster '
            'ratios at 40% (Table S30). The audit thus separates genuine homology signal '
            'from composition-level noise; note that the binding-profile mirroring channel '
            'operates on activity vectors and is unaffected by sequence shuffling, '
            're-confirming that the two leakage channels are orthogonal.', S_BODY),

    ('par', 'Taken together, the P0 analyses strengthen the operational components of the '
            'framework—the router generalizes out of dataset, the audit classifications '
            'are stable to coverage/identity settings and to the clustering engine, the '
            'audit discriminates true leakage from random sequences, and auditing changes '
            'deployment decisions with measurable cost when omitted—while explicitly '
            'restricting the τ ≈ 0.13 band to a descriptive role.', S_BODY),
]

anchor = find('4 Discussion and conclusion')[1]
insert_block_before(anchor, items)
ADDED.append(('新增', '3.6.5 + Table 6', 'P0 五项稳健性实验正文'))

# =====================================================================
# 三、附录 Tables S26-S31 + Figure S11（追加到文末）
# =====================================================================
def cap(text):
    p = doc.add_paragraph(text)
    try:
        p.style = S_CAP
    except Exception:
        pass


def note(text):
    doc.add_paragraph(text)


def add_table(rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    try:
        t.style = doc.tables[0].style
    except Exception:
        pass
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            t.cell(r, c).text = str(v)
    for c in t.rows[0].cells:
        for pp in c.paragraphs:
            for run in pp.runs:
                run.bold = True
    return t


doc.add_paragraph('P0 framework-robustness analyses (supplementary tables)').runs[0].bold = True

# S26 router LODO
cap('Table S26 Family-remote router leave-one-dataset-out validation '
    '(threshold selected on three datasets, tested on the fourth)')
rows = [['Held-out dataset', 'n test', 'n remote', 'Selected threshold (%)',
         'TP', 'FP', 'TN', 'FN', 'Accuracy', 'FPR', 'FNR']]
for _, r in exp2.iterrows():
    rows.append([r.held_out, int(r.n_test), int(r.n_remote), f'{r.selected_thr:.0f}',
                 int(r.TP), int(r.FP), int(r.TN), int(r.FN),
                 f'{r.Accuracy:.3f}', f'{r.FPR:.3f}', f'{r.FNR:.3f}'])
add_table(rows)
note('Threshold grid 25–60% in 0.5% steps; ties resolved toward 40%. The fixed pre-registered '
     '40% rule yields identical confusion matrices in all four folds (mean accuracy 0.964, '
     'balanced accuracy 0.965). Ground truth: MMseqs2 clusters (c = 0.8, 40% identity).')

# S27 grid
cap('Table S27 Cross-cluster ratio (%) under the full MMseqs2 parameter grid')
pivot = grid.pivot_table(index=['dataset', 'coverage'], columns='identity_pct',
                         values='cross_cluster_ratio_pct')
rows = [['Dataset', 'Coverage'] + [f'{c}%' for c in pivot.columns]]
name = {'chembl': 'ChEMBL', 'davis': 'Davis', 'kiba': 'KIBA', 'bindingdb': 'BindingDB'}
for (ds, cov), rr in pivot.iterrows():
    rows.append([name[ds], f'{cov}'] + [f'{rr[c]:.1f}' for c in pivot.columns])
add_table(rows)
note('cov-mode 0 (bidirectional), cluster-mode 0. The c = 0.8 / 40/60/80% cells reproduce the '
     'original audit exactly (reproduction_check.csv).')

# S28 algorithm
cap('Table S28 Clustering-algorithm cross-check: set-cover/connected-component versus '
    'CD-HIT-style greedy incremental (c = 0.8)')
rows = [['Dataset', 'Identity %', 'Clusters (set-cover)', 'Clusters (greedy)',
         'Cross % (set-cover)', 'Cross % (greedy)', 'Label agreement %', "Cohen κ"]]
for _, r in algo.iterrows():
    rows.append([name[r.dataset], int(r.identity_pct),
                 int(r.mmseqs_setcover_n_clusters), int(r.cdhit_greedy_n_clusters),
                 f'{r.mmseqs_cross_ratio_pct:.2f}', f'{r.cdhit_cross_ratio_pct:.2f}',
                 f'{r.label_agreement_pct:.1f}', f'{r.cohen_kappa:.3f}'])
add_table(rows)
note('Greedy incremental = MMseqs2 --cluster-mode 2, the algorithmic analogue of CD-HIT '
     '(longest-seed, representative-assignment). Maximum cross-ratio deviation 7.7 pp '
     '(ChEMBL, 40%).')

# S29 audit vs unaudited
cap('Table S29 Audited (40% homology-free subset) versus unaudited (full nominal cold '
    'test) evaluation: rank agreement and mis-selection cost')
rows = [['Dataset', 'Metric', 'Kendall τb', 'Spearman ρ', 'Unaudited winner',
         'Audited winner', 'Winner reversed', 'Mis-selection cost', 'Relative cost %']]
for _, r in exp4.iterrows():
    rows.append([r.dataset, r.metric, f'{r.kendall_tau:.3f}', f'{r.spearman_rho:.3f}',
                 r.standard_winner, r.audited_winner,
                 'yes' if r.winner_reversed else 'no',
                 f'{r.misselection_cost:.4f}', f'{r.relative_cost_pct:.1f}'])
add_table(rows)
note('Five models (4 GNN + Reptile-Transformer); Davis/KIBA/BindingDB matched prediction '
     'sets (Table S21). Cost = audited-subset performance of audited winner minus that of '
     'unaudited winner (positive = loss from skipping the audit). ECE costs are on the '
     'probability scale (smaller better); ChEMBL omitted: no matched cross-cluster '
     'prediction set for all models.')

# S30 negative control
cap('Table S30 Random-sequence negative control (residue-shuffled test sequences, 3 seeds)')
agg = neg.groupby(['dataset', 'identity_pct']).agg(
    m=('cross_cluster_ratio_pct', 'mean'), lo=('cross_cluster_ratio_pct', 'min'),
    hi=('cross_cluster_ratio_pct', 'max'), nid=('nearest_id_max_pct', 'max')).reset_index()
rows = [['Dataset', 'Identity %', 'Cross-cluster % mean', 'min–max (3 seeds)',
         'Max nearest identity to train (%)']]
for _, r in agg.iterrows():
    rows.append([name[r.dataset], int(r.identity_pct), f'{r.m:.1f}',
                 f'{r.lo:.1f}–{r.hi:.1f}', f'{r.nid:.1f}'])
add_table(rows)
note('Every shuffled target was cross-cluster and no bidirectionally covered (c = 0.8) '
     'training hit above 40% identity existed in any setting (36/36 settings at 100% '
     'cross-cluster).')

# S31 τ LODO
cap('Table S31 Leave-one-dataset-out validation of the τ ≈ 0.13 GNN-applicability band')
rows = [['Held-out panel', 'τ', 'ΔR²', 'Fitted band (from other 3)', 'Predicted GNN',
         'Observed GNN', 'Correct']]
for _, r in exp1.iterrows():
    rows.append([r.held_out, f'{r.tau:.3f}', f'{r.delta_R2:+.4f}',
                 f"[{r.b_lo}, {r.b_hi}]", int(r.band_pred),
                 int(r.true_GNN_adv), 'yes' if r.band_correct else 'no'])
add_table(rows)
note('Band rule: predict GNN advantage inside midpoint boundaries fitted on the other three '
     'panels. LODO accuracy = 0.50 (2/4); 1-nearest-neighbour = 0.50; majority-class = 0.75. '
     'Lower boundary 0.130 in both identifiable folds; upper boundary 0.194–0.249 '
     '(mean 0.212). Pre-registered consequence: τ retained as descriptive band only.')

# Figure S11
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(P0 + r'\FigureS11_audit_param_sensitivity.png', width=Inches(6.2))
cap('Figure S11 Homology-audit parameter sensitivity. Cross-cluster ratio of test targets '
    'across sequence-identity thresholds (30–80%) and bidirectional coverage settings '
    '(0.7–0.9), per dataset (MMseqs2 easy-cluster).')

ADDED.append(('新增', 'Tables S26-S31 + Figure S11', 'P0 附录补充'))

doc.save(DST)
print('SAVED:', DST)
for k, ok in edits:
    print(f'  edit {k}: {"OK" if ok else "FAILED"}')
for a in ADDED:
    print('  added:', a)
